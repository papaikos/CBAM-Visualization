from __future__ import annotations

import unittest
import zipfile
from io import BytesIO

from openpyxl import Workbook, load_workbook

from batch_excel import (
    MAX_UNCOMPRESSED_BYTES,
    MAX_WORKBOOK_BYTES,
    WorkbookFormatError,
    WorkbookTooLarge,
    build_input_template,
    build_report_workbook,
    parse_input_workbook,
    safe_excel_text,
    validate_xlsx_archive,
)
from batch_report import AGGREGATED_DUPLICATE_WARNING


def workbook_bytes(headers: list[str], rows: list[list[object]], sheet_name: str = "Input") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class BatchExcelTests(unittest.TestCase):
    def test_template_has_professional_input_contract_and_hidden_references(self) -> None:
        data = build_input_template(
            {
                "codes": ["25231000", "76011010"],
                "countries": ["Egypt", "Türkiye"],
                "years": [2026, 2027],
            }
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Instructions", "Input", "Reference"])
        self.assertEqual(workbook["Reference"].sheet_state, "hidden")
        input_sheet = workbook["Input"]
        self.assertEqual(
            [cell.value for cell in input_sheet[1]],
            ["CN Code", "Country", "Weight (tonnes)", "CO2 Price (EUR/tCO2)"],
        )
        self.assertEqual(input_sheet.freeze_panes, "A2")
        self.assertEqual(input_sheet.auto_filter.ref, "A1:D501")
        self.assertEqual(input_sheet["A2"].number_format, "00000000")
        self.assertEqual(workbook["Reference"]["A2"].value, 25231000)
        self.assertEqual(workbook["Reference"]["A2"].number_format, "00000000")
        self.assertEqual(len(input_sheet.data_validations.dataValidation), 2)
        self.assertEqual(workbook["Instructions"].print_area, "'Instructions'!$A$1:$D$14")
        self.assertEqual(workbook["Instructions"].page_setup.fitToWidth, 1)
        self.assertEqual(workbook["Instructions"].page_setup.orientation, "landscape")
        self.assertGreaterEqual(
            workbook["Instructions"].row_dimensions[1].height or 0,
            30,
        )
        self.assertEqual(input_sheet.print_area, "'Input'!$A$1:$D$25")
        self.assertEqual(input_sheet.page_setup.fitToWidth, 1)
        self.assertEqual(workbook.defined_names["ValidCNCodes"].attr_text, "'Reference'!$A$2:$A$3")
        self.assertEqual(workbook.defined_names["ValidCountries"].attr_text, "'Reference'!$B$2:$B$3")

    def test_import_accepts_trimmed_case_insensitive_headers_and_defaults_weight(self) -> None:
        data = workbook_bytes(
            [" cn code ", "COUNTRY", "Weight (tonnes)", "co2 price (eur/tco2)"],
            [["076011010", "Türkiye", None, 65], [25231000, "Egypt", 12.5, None]],
        )
        parsed = parse_input_workbook(data)
        self.assertEqual(
            parsed,
            {
                "lines": [
                    {
                        "inputLine": 2,
                        "cnCode": "076011010",
                        "country": "Türkiye",
                        "weightTonnes": 1.0,
                        "co2Price": 65.0,
                    },
                    {
                        "inputLine": 3,
                        "cnCode": "25231000",
                        "country": "Egypt",
                        "weightTonnes": 12.5,
                        "co2Price": None,
                    },
                ],
                "issues": [],
            },
        )

    def test_import_restores_eight_digit_cn_code_from_numeric_excel_cell(self) -> None:
        data = workbook_bytes(
            ["CN Code", "Country", "Weight (tonnes)", "CO2 Price (EUR/tCO2)"],
            [[1012100, "Greece", 1, None]],
        )
        parsed = parse_input_workbook(data)
        self.assertEqual(parsed["lines"][0]["cnCode"], "01012100")

    def test_import_accepts_decimal_comma_text_from_localized_excel_cells(self) -> None:
        data = workbook_bytes(
            ["CN Code", "Country", "Weight (tonnes)", "CO2 Price (EUR/tCO2)"],
            [[76011010, "Greece", "1,5", "69,00"]],
        )
        parsed = parse_input_workbook(data)
        self.assertEqual(parsed["issues"], [])
        self.assertEqual(parsed["lines"][0]["weightTonnes"], 1.5)
        self.assertEqual(parsed["lines"][0]["co2Price"], 69.0)

    def test_import_rejects_formula_cells_instead_of_silently_using_defaults(self) -> None:
        data = workbook_bytes(
            ["CN Code", "Country", "Weight (tonnes)", "CO2 Price (EUR/tCO2)"],
            [[76011010, "Greece", "=1+1", "=60+5"]],
        )
        parsed = parse_input_workbook(data)
        self.assertEqual(parsed["lines"], [])
        self.assertEqual(
            parsed["issues"][0]["error"],
            "Excel formulas are not supported. Paste calculated values instead.",
        )

    def test_import_finds_the_input_worksheet_case_insensitively(self) -> None:
        data = workbook_bytes(
            ["CN Code", "Country", "Weight (tonnes)", "CO2 Price (EUR/tCO2)"],
            [[76011010, "Greece", 1, 65]],
            sheet_name="INPUT",
        )
        parsed = parse_input_workbook(data)
        self.assertEqual(len(parsed["lines"]), 1)
        self.assertEqual(parsed["issues"], [])

    def test_import_keeps_valid_rows_and_reports_partially_completed_rows(self) -> None:
        data = workbook_bytes(
            ["CN Code", "Country", "Weight (tonnes)", "CO2 Price (EUR/tCO2)"],
            [["76011010", "Türkiye", 5, 70], ["25231000", None, 1, None], [None, None, None, None]],
        )
        parsed = parse_input_workbook(data)
        self.assertEqual(len(parsed["lines"]), 1)
        self.assertEqual(parsed["issues"][0]["inputLine"], 3)
        self.assertEqual(parsed["issues"][0]["error"], "Country is required.")

    def test_missing_input_or_invalid_headers_rejects_workbook(self) -> None:
        with self.assertRaisesRegex(WorkbookFormatError, "Input worksheet"):
            parse_input_workbook(workbook_bytes(["CN Code"], [], sheet_name="Data"))
        with self.assertRaisesRegex(WorkbookFormatError, "Unknown header"):
            parse_input_workbook(
                workbook_bytes(
                    ["CN Code", "Country", "Weight (tonnes)", "Surprise"], []
                )
            )
        with self.assertRaisesRegex(WorkbookFormatError, "Duplicate header"):
            parse_input_workbook(
                workbook_bytes(
                    ["CN Code", "Country", "Weight (tonnes)", "country"], []
                )
            )

    def test_archive_safeguards_reject_compressed_and_uncompressed_limits(self) -> None:
        with self.assertRaises(WorkbookTooLarge):
            validate_xlsx_archive(b"x" * (MAX_WORKBOOK_BYTES + 1))

        archive = BytesIO()
        with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zipped:
            zipped.writestr("xl/worksheets/sheet1.xml", b"0" * (MAX_UNCOMPRESSED_BYTES + 1))
        with self.assertRaisesRegex(WorkbookTooLarge, "uncompressed"):
            validate_xlsx_archive(archive.getvalue())

    def test_safe_excel_text_prevents_formula_injection(self) -> None:
        self.assertEqual(safe_excel_text("=1+1"), "'=1+1")
        self.assertEqual(safe_excel_text("+SUM(A1:A2)"), "'+SUM(A1:A2)")
        self.assertEqual(safe_excel_text("ordinary"), "ordinary")

    def test_report_workbook_has_numeric_outputs_blanks_and_literal_text(self) -> None:
        report = {
            "year": 2026,
            "dataMode": "aggregated_sqlite",
            "results": [
                {
                    "inputLine": 1,
                    "sourceRowId": 8,
                    "sourceLine": 8,
                    "cnCode": "76011010",
                    "country": "=1+1",
                    "year": 2026,
                    "sourceProductionRoute": "@route",
                    "productionRouteLabel": "@route",
                    "weightTonnes": 1.0,
                    "emissionsPerTonne": 1.2,
                    "totalEmissions": None,
                    "co2Price": None,
                    "costPerTonne": None,
                    "totalCbamCost": None,
                    "duplicateGroupSize": 2,
                    "hasDuplicateSourceRows": True,
                }
            ],
            "issues": [
                {
                    "inputLine": 2,
                    "cnCode": "25231000",
                    "country": "Egypt",
                    "weightTonnes": "=1+1",
                    "co2Price": "+2",
                    "error": "Weight must be a number greater than zero.",
                }
            ],
            "duplicateWarnings": [
                {"inputLine": 1, "message": AGGREGATED_DUPLICATE_WARNING}
            ],
        }
        data = build_report_workbook(report)
        workbook = load_workbook(BytesIO(data), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["Report", "Input Issues", "Methodology & Disclaimer"],
        )
        sheet = workbook["Report"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        self.assertEqual(sheet.cell(2, headers["Country"]).value, "'=1+1")
        self.assertEqual(sheet.cell(2, headers["Country"]).data_type, "s")
        self.assertEqual(sheet.cell(2, headers["Source Route Value"]).value, "'@route")
        self.assertIsInstance(
            sheet.cell(2, headers["Weight (tonnes)"]).value,
            (int, float),
        )
        self.assertIsNone(sheet.cell(2, headers["Total Emissions (tCO2)"]).value)
        self.assertIsNone(sheet.cell(2, headers["CO2 Price (EUR/tCO2)"]).value)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.auto_filter.ref)
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertEqual(workbook["Input Issues"].page_setup.fitToWidth, 1)
        self.assertEqual(
            workbook["Methodology & Disclaimer"].page_setup.orientation,
            "portrait",
        )
        self.assertGreaterEqual(
            workbook["Methodology & Disclaimer"].row_dimensions[1].height or 0,
            30,
        )
        self.assertFalse(
            workbook["Methodology & Disclaimer"].print_options.horizontalCentered
        )
        self.assertLessEqual(
            workbook["Methodology & Disclaimer"].column_dimensions["A"].width,
            92,
        )
        self.assertGreaterEqual(
            workbook["Methodology & Disclaimer"].row_dimensions[9].height,
            80,
        )
        issues_sheet = workbook["Input Issues"]
        issue_headers = {cell.value: cell.column for cell in issues_sheet[1]}
        self.assertEqual(
            issues_sheet.cell(2, issue_headers["Weight (tonnes)"]).value,
            "'=1+1",
        )
        self.assertEqual(
            issues_sheet.cell(2, issue_headers["CO2 Price (EUR/tCO2)"]).value,
            "'+2",
        )
        self.assertIn(
            "bundled SQLite database stores one averaged record",
            workbook["Methodology & Disclaimer"]["A8"].value,
        )


if __name__ == "__main__":
    unittest.main()
