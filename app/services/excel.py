"""Excel workbook building and parsing (input template, uploads, report export)."""

from __future__ import annotations

import math
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import defusedxml  # noqa: F401 - ensures hardened XML parsers are installed for openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from app.config import ELECTRICITY_CN_CODE
from app.errors import WorkbookFormatError, WorkbookTooLarge
from app.services.batch import AGGREGATED_DUPLICATE_WARNING


ELECTRICITY_UNIT_NOTE = (
    f"Electricity (CN {ELECTRICITY_CN_CODE}) is measured per MWh: enter MWh in the weight "
    "column; its emission and cost values are per MWh instead of per tonne."
)

MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

NAVY = "0F172A"
NAVY_LIGHT = "1E293B"
TEAL = "0F766E"
TEAL_LIGHT = "CCFBF1"
SLATE = "475569"
PALE = "F8FAFC"
WHITE = "FFFFFF"
AMBER = "B45309"
AMBER_LIGHT = "FEF3C7"
THIN_BORDER = Border(bottom=Side(style="thin", color="D7E0E9"))

INPUT_HEADERS = [
    "CN Code",
    "Country",
    "Weight (tonnes)",
    "CO2 Price (EUR/tCO2)",
]
REPORT_HEADERS = [
    "Input Line",
    "Source Line",
    "Year",
    "CN Code",
    "Country",
    "Production Route",
    "Source Route Value",
    "Weight (tonnes)",
    "Emissions (tCO2/tonne)",
    "Total Emissions (tCO2)",
    "CO2 Price (EUR/tCO2)",
    "Estimated CBAM Cost (EUR/tonne)",
    "Estimated Total CBAM Cost (EUR)",
    "Duplicate Source Rows",
    "Warning",
]


def validate_xlsx_archive(data: bytes) -> None:
    if len(data) > MAX_WORKBOOK_BYTES:
        raise WorkbookTooLarge(
            "Workbook exceeds the 5 MiB upload limit.", code="workbook_too_large"
        )
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if any(member.flag_bits & 0x1 for member in members):
                raise WorkbookFormatError("Password-protected workbooks are not supported.")
            uncompressed_size = sum(member.file_size for member in members)
            if uncompressed_size > MAX_UNCOMPRESSED_BYTES:
                raise WorkbookTooLarge(
                    "Workbook uncompressed contents exceed the 50 MiB safety limit.",
                    code="workbook_uncompressed_too_large",
                )
            names = {member.filename for member in members}
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise WorkbookFormatError("The uploaded file is not a valid .xlsx workbook.")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise WorkbookFormatError("Macro-enabled workbooks are not supported.")
    except zipfile.BadZipFile as exc:
        raise WorkbookFormatError("The uploaded file is not a valid .xlsx workbook.") from exc


def safe_excel_text(value: object) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _safe_issue_value(value: object) -> object:
    if value is None or isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return safe_excel_text(value)


def _style_header(sheet, headers: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=TEAL))
    sheet.row_dimensions[1].height = 30


def _finish_table(sheet, widths: list[float], last_row: int) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{max(last_row, 2)}"
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=max(last_row, 2)):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == len(widths))


def _configure_print(
    sheet,
    print_area: str,
    *,
    orientation: str,
    one_page_tall: bool = False,
    paper_size: str = "9",
    repeat_header: bool = False,
    horizontal_centered: bool = True,
) -> None:
    sheet.print_area = print_area
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(
        fitToPage=True,
        autoPageBreaks=False,
    )
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = paper_size
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1 if one_page_tall else 0
    sheet.page_margins.left = 0.3
    sheet.page_margins.right = 0.3
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5
    sheet.print_options.horizontalCentered = horizontal_centered
    if repeat_header:
        sheet.print_title_rows = "1:1"


def build_input_template(meta: dict) -> bytes:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.sheet_view.showGridLines = False
    instructions["A1"] = "CBAM Batch Report Input"
    instructions["A1"].font = Font(size=22, bold=True, color=NAVY)
    instructions.row_dimensions[1].height = 34
    instructions["A2"] = "Prepare multiple CN code and country lookups in one workbook."
    instructions["A2"].font = Font(size=11, color=SLATE)
    instructions["A4"] = "Required columns"
    instructions["A4"].font = Font(size=12, bold=True, color=TEAL)
    for row, header in enumerate(INPUT_HEADERS, start=5):
        instructions.cell(row=row, column=1, value=header)
    instructions["C4"] = "Four-step guide"
    instructions["C4"].font = Font(size=12, bold=True, color=TEAL)
    steps = [
        "1. Open the Input worksheet.",
        "2. Choose a CN code and country for each row.",
        "3. Add an optional weight and CO2 price.",
        "4. Save as .xlsx and upload it on the Batch Report page.",
    ]
    for row, step in enumerate(steps, start=5):
        instructions.cell(row=row, column=3, value=step)
    instructions["A11"] = "Weight defaults to 1 tonne when blank. CO2 price is optional."
    instructions["A12"] = (
        "Select the reporting year on the website; it is intentionally absent from this workbook."
    )
    instructions["A13"] = ELECTRICITY_UNIT_NOTE
    instructions["A14"] = (
        "Academic estimate only — independently verify source values before compliance use."
    )
    instructions["A14"].fill = PatternFill("solid", fgColor=AMBER_LIGHT)
    instructions["A14"].font = Font(color=AMBER, bold=True)
    instructions.merge_cells("A14:D14")
    instructions.merge_cells("A1:D1")
    instructions.merge_cells("A2:D2")
    instructions.merge_cells("A11:D11")
    instructions.merge_cells("A12:D12")
    instructions.merge_cells("A13:D13")
    for cell in (
        instructions["A2"],
        instructions["A11"],
        instructions["A12"],
        instructions["A13"],
        instructions["A14"],
    ):
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    instructions.row_dimensions[2].height = 26
    instructions.row_dimensions[11].height = 28
    instructions.row_dimensions[12].height = 30
    instructions.row_dimensions[13].height = 30
    instructions.row_dimensions[14].height = 42
    instructions.column_dimensions["A"].width = 34
    instructions.column_dimensions["B"].width = 4
    instructions.column_dimensions["C"].width = 62
    instructions.column_dimensions["D"].width = 12
    _configure_print(
        instructions,
        "A1:D14",
        orientation="landscape",
        one_page_tall=True,
    )

    input_sheet = workbook.create_sheet("Input")
    _style_header(input_sheet, INPUT_HEADERS)
    _finish_table(input_sheet, [20, 30, 20, 27], 501)
    for row in range(2, 502):
        input_sheet.cell(row=row, column=1).number_format = "00000000"
        input_sheet.cell(row=row, column=3).number_format = "#,##0.00"
        input_sheet.cell(row=row, column=4).number_format = "€#,##0.00"
    _configure_print(
        input_sheet,
        "A1:D25",
        orientation="landscape",
        repeat_header=True,
    )

    reference = workbook.create_sheet("Reference")
    reference.append(["Valid CN Codes", "Valid Countries"])
    codes = [str(value) for value in meta.get("codes", [])]
    countries = [str(value) for value in meta.get("countries", [])]
    for row_index in range(max(len(codes), len(countries))):
        code = codes[row_index] if row_index < len(codes) else None
        reference.append(
            [
                int(code) if code and code.isdigit() else code,
                countries[row_index] if row_index < len(countries) else None,
            ]
        )
        reference.cell(row=row_index + 2, column=1).number_format = "00000000"
    _style_header(reference, ["Valid CN Codes", "Valid Countries"])
    reference.column_dimensions["A"].width = 22
    reference.column_dimensions["B"].width = 32
    reference.sheet_state = "hidden"

    code_end = max(2, len(codes) + 1)
    country_end = max(2, len(countries) + 1)
    workbook.defined_names.add(
        DefinedName("ValidCNCodes", attr_text=f"'Reference'!$A$2:$A${code_end}")
    )
    workbook.defined_names.add(
        DefinedName("ValidCountries", attr_text=f"'Reference'!$B$2:$B${country_end}")
    )
    code_validation = DataValidation(type="list", formula1="=ValidCNCodes", allow_blank=True)
    country_validation = DataValidation(
        type="list", formula1="=ValidCountries", allow_blank=True
    )
    input_sheet.add_data_validation(code_validation)
    input_sheet.add_data_validation(country_validation)
    code_validation.add("A2:A501")
    country_validation.add("B2:B501")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _normalized_header(value: Any) -> str:
    return "" if value is None else str(value).strip().casefold()


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cn_identifier(value: Any) -> str:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return f"{value:08d}"
    if isinstance(value, float) and value.is_integer():
        return f"{int(value):08d}"
    return _identifier(value)


def _optional_number(value: Any, *, default: float | None, message: str) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        raise ValueError(message)
    if isinstance(value, str):
        candidate = value.strip()
        if candidate.count(",") == 1 and "." not in candidate:
            value = candidate.replace(",", ".")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(message) from exc
    if not math.isfinite(number):
        raise ValueError(message)
    return number


def parse_input_workbook(data: bytes) -> dict:
    validate_xlsx_archive(data)
    try:
        workbook = load_workbook(
            BytesIO(data),
            read_only=True,
            data_only=False,
            keep_vba=False,
            keep_links=False,
        )
    except Exception as exc:
        raise WorkbookFormatError("The uploaded file could not be read as an .xlsx workbook.") from exc
    try:
        input_sheet_name = next(
            (name for name in workbook.sheetnames if name.strip().casefold() == "input"),
            None,
        )
        if input_sheet_name is None:
            raise WorkbookFormatError("The workbook must contain an Input worksheet.")
        sheet = workbook[input_sheet_name]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        normalized = [_normalized_header(value) for value in first_row]
        expected = {header.casefold(): header for header in INPUT_HEADERS}
        nonblank = [value for value in normalized if value]
        duplicates = sorted({value for value in nonblank if nonblank.count(value) > 1})
        if duplicates:
            raise WorkbookFormatError(f"Duplicate header: {duplicates[0]}")
        unknown = [value for value in nonblank if value not in expected]
        if unknown:
            raise WorkbookFormatError(f"Unknown header: {unknown[0]}")
        missing = [header for key, header in expected.items() if key not in nonblank]
        if missing:
            raise WorkbookFormatError(f"Missing header: {missing[0]}")
        positions = {value: index for index, value in enumerate(normalized)}

        lines: list[dict] = []
        issues: list[dict] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            def read(header: str) -> Any:
                index = positions[header.casefold()]
                return values[index] if index < len(values) else None

            raw_code = read("CN Code")
            raw_country = read("Country")
            raw_weight = read("Weight (tonnes)")
            raw_price = read("CO2 Price (EUR/tCO2)")
            if all(
                value is None or (isinstance(value, str) and not value.strip())
                for value in (raw_code, raw_country, raw_weight, raw_price)
            ):
                continue
            code = _cn_identifier(raw_code)
            country = _identifier(raw_country)
            base = {
                "inputLine": row_number,
                "cnCode": code,
                "country": country,
                "weightTonnes": raw_weight,
                "co2Price": raw_price,
            }
            if any(
                isinstance(value, str) and value.lstrip().startswith("=")
                for value in (raw_code, raw_country, raw_weight, raw_price)
            ):
                issues.append(
                    {
                        **base,
                        "error": (
                            "Excel formulas are not supported. "
                            "Paste calculated values instead."
                        ),
                    }
                )
                continue
            if not code:
                issues.append({**base, "error": "CN code is required."})
                continue
            if not country:
                issues.append({**base, "error": "Country is required."})
                continue
            try:
                weight = _optional_number(
                    raw_weight,
                    default=1.0,
                    message="Weight must be a number greater than zero.",
                )
                if weight is None or weight <= 0:
                    raise ValueError("Weight must be a number greater than zero.")
            except ValueError as exc:
                issues.append({**base, "error": str(exc)})
                continue
            try:
                price = _optional_number(
                    raw_price,
                    default=None,
                    message="CO₂ price must be a number greater than or equal to zero.",
                )
                if price is not None and price < 0:
                    raise ValueError(
                        "CO₂ price must be a number greater than or equal to zero."
                    )
            except ValueError as exc:
                issues.append({**base, "error": str(exc)})
                continue
            lines.append(
                {
                    "inputLine": row_number,
                    "cnCode": code,
                    "country": country,
                    "weightTonnes": weight,
                    "co2Price": price,
                }
            )
        return {"lines": lines, "issues": issues}
    finally:
        workbook.close()


def build_report_workbook(report: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    _style_header(sheet, REPORT_HEADERS)
    for item in report.get("results", []):
        warning = AGGREGATED_DUPLICATE_WARNING if item.get("hasDuplicateSourceRows") else ""
        sheet.append(
            [
                item.get("inputLine"),
                item.get("sourceLine"),
                item.get("year"),
                safe_excel_text(item.get("cnCode")),
                safe_excel_text(item.get("country")),
                safe_excel_text(item.get("productionRouteLabel")),
                safe_excel_text(item.get("sourceProductionRoute")),
                item.get("weightTonnes"),
                item.get("emissionsPerTonne"),
                item.get("totalEmissions"),
                item.get("co2Price"),
                item.get("costPerTonne"),
                item.get("totalCbamCost"),
                item.get("duplicateGroupSize"),
                safe_excel_text(warning),
            ]
        )
    _finish_table(
        sheet,
        [12, 13, 10, 16, 24, 22, 22, 17, 23, 23, 22, 30, 31, 22, 72],
        sheet.max_row,
    )
    _configure_print(
        sheet,
        f"A1:O{max(sheet.max_row, 2)}",
        orientation="landscape",
        paper_size="8",
        repeat_header=True,
    )
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 4).number_format = "@"
        for column in (8, 9, 10):
            sheet.cell(row, column).number_format = "#,##0.0000"
        for column in (11, 12, 13):
            sheet.cell(row, column).number_format = "€#,##0.00"

    issues_sheet = workbook.create_sheet("Input Issues")
    issue_headers = [
        "Input Line",
        "CN Code",
        "Country",
        "Weight (tonnes)",
        "CO2 Price (EUR/tCO2)",
        "Error",
    ]
    _style_header(issues_sheet, issue_headers)
    for issue in report.get("issues", []):
        issues_sheet.append(
            [
                issue.get("inputLine"),
                safe_excel_text(issue.get("cnCode")),
                safe_excel_text(issue.get("country")),
                _safe_issue_value(issue.get("weightTonnes")),
                _safe_issue_value(issue.get("co2Price")),
                safe_excel_text(issue.get("error")),
            ]
        )
    _finish_table(issues_sheet, [12, 18, 24, 20, 24, 58], issues_sheet.max_row)
    _configure_print(
        issues_sheet,
        f"A1:F{max(issues_sheet.max_row, 2)}",
        orientation="landscape",
        repeat_header=True,
    )

    methodology = workbook.create_sheet("Methodology & Disclaimer")
    methodology.sheet_view.showGridLines = False
    methodology["A1"] = f"CBAM Batch Report — {report.get('year', '')}"
    methodology["A1"].font = Font(size=22, bold=True, color=NAVY)
    methodology.row_dimensions[1].height = 34
    methodology["A3"] = "Calculation methodology"
    methodology["A3"].font = Font(size=12, bold=True, color=TEAL)
    methodology["A4"] = "Emissions per tonne = stored paid emissions"
    methodology["A5"] = "Total emissions = emissions per tonne × weight tonnes"
    methodology["A6"] = "Estimated CBAM cost = emissions × supplied CO2 price"
    methodology["A7"] = (
        "Conditional totals remain blank when weight equals 1 tonne or no price is supplied."
    )
    methodology["A8"] = (
        "Source limitation: the bundled SQLite database stores one averaged record per CN "
        "code, country, year, and production route, plus duplicate_count. Individual original "
        "source values cannot be reconstructed and are not fabricated."
    )
    methodology["A9"] = AGGREGATED_DUPLICATE_WARNING
    methodology["A10"] = ELECTRICITY_UNIT_NOTE
    methodology["A11"] = (
        "This informal academic tool is for general information only and is not legal, tax, "
        "accounting, financial, regulatory, or CBAM compliance advice."
    )
    methodology["A13"] = (
        "Generated in UTC: "
        + datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    )
    methodology.column_dimensions["A"].width = 92
    for row in range(1, 14):
        methodology.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
    methodology.row_dimensions[8].height = 54
    methodology.row_dimensions[9].height = 84
    methodology.row_dimensions[10].height = 30
    methodology["A9"].fill = PatternFill("solid", fgColor=AMBER_LIGHT)
    methodology["A9"].font = Font(color=AMBER, bold=True)
    _configure_print(
        methodology,
        "A1:A13",
        orientation="portrait",
        one_page_tall=True,
        horizontal_centered=False,
    )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
